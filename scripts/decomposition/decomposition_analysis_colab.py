"""
[Colab version] Scale-dependent decomposition of exposure change (2030 -> 2060)
into Climate effect / Population effect / Interaction effect,
computed at three spatial scales: 1km grid, municipal (161), provincial (17).

Folder: /content/drive/MyDrive/Colab Notebooks/EHE/Fig4
  - 2030_pop_1km.tif, 2060_pop_1km.tif
  - heatdays_2030_v2.tif, heatdays_2060_v2.tif
  - Method_a_zonal_municipal.xlsx (Sheet3)
  - Method_a_zonal_provincial.xlsx (Sheet3)
"""

# ---------------------------------------------------------------------------
# 0. Google Drive mount + package install
# ---------------------------------------------------------------------------
from google.colab import drive
drive.mount('/content/drive')

!pip install -q tifffile openpyxl

import os
import numpy as np
import tifffile
import openpyxl

BASE_DIR = "/content/drive/MyDrive/Colab Notebooks/EHE/Fig4"


# ---------------------------------------------------------------------------
# 1. Pure-python TIFF (early-change) LZW decoder
#    (kept as a fallback in case imagecodecs isn't installed; if it IS
#     available, tifffile.imread() would also just work directly)
# ---------------------------------------------------------------------------
def lzw_decode(data: bytes) -> bytes:
    CLEAR = 256
    EOI = 257
    out = bytearray()

    def reset_table():
        return [bytes([i]) for i in range(256)] + [b"", b""]

    table = reset_table()
    bitbuf = 0
    bitcnt = 0
    pos = 0
    codelen = 9
    prev = None
    n = len(data)

    def next_code():
        nonlocal bitbuf, bitcnt, pos
        while bitcnt < codelen:
            if pos >= n:
                return None
            bitbuf = (bitbuf << 8) | data[pos]
            pos += 1
            bitcnt += 8
        bitcnt -= codelen
        return (bitbuf >> bitcnt) & ((1 << codelen) - 1)

    while True:
        code = next_code()
        if code is None or code == EOI:
            break
        if code == CLEAR:
            table = reset_table()
            codelen = 9
            prev = None
            continue
        if prev is None:
            entry = table[code]
            out += entry
            prev = entry
        else:
            if code < len(table):
                entry = table[code]
            elif code == len(table):
                entry = prev + prev[:1]
            else:
                raise ValueError(f"bad LZW code {code}, table size {len(table)}")
            out += entry
            table.append(prev + entry[:1])
            prev = entry
        # TIFF "early change": bump code length one code early
        L = len(table)
        if L >= 511 and codelen == 9:
            codelen = 10
        elif L >= 1023 and codelen == 10:
            codelen = 11
        elif L >= 2047 and codelen == 11:
            codelen = 12
    return bytes(out)


def read_tiled_lzw_f64(path: str) -> np.ndarray:
    """Read a tiled, LZW-compressed, float64, predictor=1 GeoTIFF (as exported
    by ArcGIS Pro PointToRaster) without needing imagecodecs/gdal."""
    t = tifffile.TiffFile(path)
    p = t.pages[0]
    w, h = p.tags["ImageWidth"].value, p.tags["ImageLength"].value
    tw, tl = p.tags["TileWidth"].value, p.tags["TileLength"].value
    offs = p.tags["TileOffsets"].value
    counts = p.tags["TileByteCounts"].value
    ntx = (w + tw - 1) // tw
    nty = (h + tl - 1) // tl

    arr = np.full((h, w), np.nan, dtype=np.float64)
    with open(path, "rb") as f:
        raw = f.read()

    idx = 0
    for ty in range(nty):
        for tx in range(ntx):
            off, cnt = offs[idx], counts[idx]
            dec = lzw_decode(raw[off : off + cnt])
            n_expected = tw * tl * 8
            if len(dec) < n_expected:
                dec = dec + b"\x00" * (n_expected - len(dec))
            vals = np.frombuffer(dec[:n_expected], dtype="<f8").reshape(tl, tw)
            y0, x0 = ty * tl, tx * tw
            y1, x1 = min(y0 + tl, h), min(x0 + tw, w)
            arr[y0:y1, x0:x1] = vals[0 : (y1 - y0), 0 : (x1 - x0)]
            idx += 1
    return arr


def read_pop_tif(path: str) -> np.ndarray:
    """Try the normal tifffile path first (works if imagecodecs is present,
    e.g. usually true on Colab); fall back to the pure-python LZW decoder."""
    try:
        return tifffile.imread(path).astype(np.float64)
    except Exception:
        return read_tiled_lzw_f64(path)


# ---------------------------------------------------------------------------
# 2. Decomposition core: E = P * H  =>  dE = P0*dH + H0*dP + dP*dH
# ---------------------------------------------------------------------------
def decompose(P0, P1, H0, H1):
    P0 = np.asarray(P0, dtype=float)
    P1 = np.asarray(P1, dtype=float)
    H0 = np.asarray(H0, dtype=float)
    H1 = np.asarray(H1, dtype=float)
    dP, dH = P1 - P0, H1 - H0

    climate = float((P0 * dH).sum())        # climate effect, holding P at base
    population = float((H0 * dP).sum())     # population effect, holding H at base
    interaction = float((dP * dH).sum())    # joint/interaction term
    E0, E1 = float((P0 * H0).sum()), float((P1 * H1).sum())

    return dict(
        n=len(P0), E0=E0, E1=E1, delta=E1 - E0,
        climate=climate, population=population, interaction=interaction,
        check=climate + population + interaction,  # should equal `delta`
    )


# ---------------------------------------------------------------------------
# 3. Grid level (1km): decode rasters, resample heatdays onto the pop grid
# ---------------------------------------------------------------------------
def grid_level_decomposition():
    pop30 = read_pop_tif(os.path.join(BASE_DIR, "2030_pop_1km.tif"))
    pop60 = read_pop_tif(os.path.join(BASE_DIR, "2060_pop_1km.tif"))
    pop30_mask = (pop30 > -1e30) & (pop30 < 1e15) & np.isfinite(pop30)
    pop60_mask = (pop60 > -1e30) & (pop60 < 1e15) & np.isfinite(pop60)

    # pop grid geotransform (from .tfw): upper-left x, y, cellsize
    POP_X0, POP_Y0, POP_CS = 746611.0313, 2067943.9535, 1000.0
    H, W = pop30.shape

    heat30 = tifffile.imread(os.path.join(BASE_DIR, "heatdays_2030_v2.tif")).astype(np.float64)
    heat60 = tifffile.imread(os.path.join(BASE_DIR, "heatdays_2060_v2.tif")).astype(np.float64)
    heat30_mask = (heat30 > -1e30) & (heat30 < 1e6) & np.isfinite(heat30)
    heat60_mask = (heat60 > -1e30) & (heat60 < 1e6) & np.isfinite(heat60)

    # heatdays grid geotransform (cellsize is NOT 1000m -> 1046.7390550674m)
    HEAT_X0, HEAT_Y0, HEAT_CS = 746403.5065026443, 2068142.6506664562, 1046.7390550674
    Hh, Wh = heat30.shape

    # nearest-neighbor resample heatdays onto the population grid, using
    # each population cell's center coordinate (same method used to build
    # the official EHE rasters, validated pixel-for-pixel earlier)
    rows, cols = np.arange(H), np.arange(W)
    cc, rr = np.meshgrid(cols, rows)
    x_c = POP_X0 + (cc + 0.5) * POP_CS
    y_c = POP_Y0 - (rr + 0.5) * POP_CS
    hc = np.floor((x_c - HEAT_X0) / HEAT_CS).astype(int)
    hr = np.floor((HEAT_Y0 - y_c) / HEAT_CS).astype(int)
    in_bounds = (hc >= 0) & (hc < Wh) & (hr >= 0) & (hr < Hh)
    hc_c, hr_c = np.clip(hc, 0, Wh - 1), np.clip(hr, 0, Hh - 1)

    heat30_rs = np.full((H, W), np.nan)
    heat60_rs = np.full((H, W), np.nan)
    m30 = heat30_mask[hr_c, hc_c] & in_bounds
    m60 = heat60_mask[hr_c, hc_c] & in_bounds
    heat30_rs[m30] = heat30[hr_c[m30], hc_c[m30]]
    heat60_rs[m60] = heat60[hr_c[m60], hc_c[m60]]

    final_mask = pop30_mask & pop60_mask & np.isfinite(heat30_rs) & np.isfinite(heat60_rs)
    print("valid grid cells:", final_mask.sum())

    return decompose(
        pop30[final_mask], pop60[final_mask],
        heat30_rs[final_mask], heat60_rs[final_mask],
    )


# ---------------------------------------------------------------------------
# 4. Municipal / provincial level: read pre-computed zonal stats (Method A)
# ---------------------------------------------------------------------------
def zonal_level_decomposition(xlsx_name, pop30_col, pop60_col, h30_col, h60_col, sheet="Sheet3"):
    wb = openpyxl.load_workbook(os.path.join(BASE_DIR, xlsx_name), data_only=True)
    ws = wb[sheet]
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    P0 = [r[idx[pop30_col]] for r in rows]
    P1 = [r[idx[pop60_col]] for r in rows]
    H0 = [r[idx[h30_col]] for r in rows]
    H1 = [r[idx[h60_col]] for r in rows]
    return decompose(P0, P1, H0, H1)


# ---------------------------------------------------------------------------
# 5. Run all three scales + print comparison
# ---------------------------------------------------------------------------
res_grid = grid_level_decomposition()
res_muni = zonal_level_decomposition(
    "Method_a_zonal_municipal.xlsx", "Pop30", "Pop60", "EHEmean30", "EHEmean60"
)
res_prov = zonal_level_decomposition(
    "Method_a_zonal_provincial.xlsx", "pop30", "pop60", "Heatday30", "Heatday60"
)

for name, res in [("GRID (1km)", res_grid), ("MUNICIPAL (161)", res_muni), ("PROVINCIAL (17)", res_prov)]:
    print(f"\n--- {name} ---")
    for k, v in res.items():
        print(f"  {k}: {v:,.2f}" if isinstance(v, float) else f"  {k}: {v}")

# relative bias of municipal/provincial vs the grid "ground truth"
print()
for label, val_key in [("Total change", "delta"), ("Climate", "climate"),
                        ("Population", "population"), ("Interaction", "interaction")]:
    g, m, p = res_grid[val_key], res_muni[val_key], res_prov[val_key]
    print(f"{label}: municipal bias = {(m-g)/g*100:+.1f}%,  provincial bias = {(p-g)/g*100:+.1f}%")
